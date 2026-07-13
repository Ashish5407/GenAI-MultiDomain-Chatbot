
import pandas as pd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_community.document_loaders import (
    WebBaseLoader,
    PyPDFLoader,
    TextLoader,
    UnstructuredImageLoader,
    YoutubeLoader
)
from utils.logger import logger


def load_source(source):

    try:
        # Website
        if source.startswith("http") and "github.com" not in source and "youtube.com" not in source:
            return WebBaseLoader(source).load()

        # YouTube
        elif "youtube.com" in source:
            return YoutubeLoader.from_youtube_url(source, add_video_info=True).load()

        # PDF
        elif source.endswith(".pdf"):
            return PyPDFLoader(source).load()

        # TXT
        elif source.endswith(".txt"):
            return TextLoader(source).load()

        # DOCX
        elif source.endswith(".docx"):
            docx_file = DocxDocument(source)
            text = "\n".join(paragraph.text for paragraph in docx_file.paragraphs)
            return [Document(page_content=text, metadata={"source": source})]

        # CSV
        elif source.endswith(".csv"):
            dataframe = pd.read_csv(source)
            text = dataframe.to_string(index=False)
            return [Document(page_content=text, metadata={"source": source})]

        # Excel
        elif source.endswith(".xlsx"):
            dataframe = pd.read_excel(source)
            text = dataframe.to_string(index=False)
            return [Document(page_content=text, metadata={"source": source})]

        # Image
        elif source.endswith((".png", ".jpg", ".jpeg")):
            return UnstructuredImageLoader(source).load()

        logger.warning(f"Unsupported source type: {source}")
        return []

    except Exception as e:
        logger.error(f"Failed to load source '{source}': {e}")
        return []
